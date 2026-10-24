"""
Equipment History Lookup Tool for Karyalaya AI.

Searches all finalized .docx deliverables in workspace/deliverables/ for a given
equipment tag (e.g. CP-402-B). Text is extracted via python-docx first — raw .docx
bytes are ZIP-compressed XML and will not yield readable text matches.

Returns a list of prior inspection excerpts referencing the same tag, with filenames
and surrounding paragraph context. Used by the agent to detect reading changes over
time and surface contradictions for LLM review.
"""

import os
import re
from typing import Dict, Any, List

WORKSPACE_DIR    = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "workspace")
)
DELIVERABLES_DIR = os.path.join(WORKSPACE_DIR, "deliverables")

_EQUIP_TAG_RE = re.compile(r'\b[A-Z]{1,3}-\d{2,4}-?[A-Z]?\b')


def _extract_docx_text(file_path: str) -> str:
    """
    Extracts all paragraph text from a .docx file via python-docx.

    Raw .docx bytes are ZIP-compressed XML — regex against the file object will not
    produce text matches. This function opens the document properly and joins paragraph
    text. Tables are also included (via cell iteration) to capture tabular readings.
    """
    try:
        from docx import Document
        doc = Document(file_path)
        parts: List[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)
    except Exception as e:
        return ""


def _find_tag_excerpts(text: str, tag: str) -> List[str]:
    """
    Returns paragraphs (lines) from text that contain the equipment tag.
    Matches whole-word occurrences only (the tag must appear as a standalone token).
    """
    tag_pattern = re.compile(r'\b' + re.escape(tag) + r'\b')
    excerpts = []
    for line in text.split("\n"):
        if tag_pattern.search(line):
            excerpt = line.strip()
            if excerpt:
                excerpts.append(excerpt)
    return excerpts


def lookup_equipment_history(tag: str) -> Dict[str, Any]:
    """
    Searches all finalized .docx deliverables for mentions of the given equipment tag.

    Skips DRAFT_ files and _trace.md files — only searches finalized official documents.
    Extracts text via python-docx, then matches with regex.

    Args:
        tag: Equipment tag string, e.g. 'CP-402-B'.

    Returns:
        Dict with status, tag, match_count, matches (list of {filename, excerpts}).
    """
    if not os.path.exists(DELIVERABLES_DIR):
        return {
            "status":      "warning",
            "tag":         tag,
            "match_count": 0,
            "matches":     [],
            "message":     "Deliverables directory does not exist yet.",
        }

    docx_files = [
        f for f in os.listdir(DELIVERABLES_DIR)
        if f.endswith(".docx")
        and not f.startswith("DRAFT_")
        and not f.endswith("_trace.md")
    ]

    if not docx_files:
        return {
            "status":      "ok",
            "tag":         tag,
            "match_count": 0,
            "matches":     [],
            "message":     "No finalized deliverables to search yet.",
        }

    matches: List[Dict[str, Any]] = []
    for filename in sorted(docx_files):
        file_path = os.path.join(DELIVERABLES_DIR, filename)
        text      = _extract_docx_text(file_path)
        if not text:
            continue
        excerpts = _find_tag_excerpts(text, tag)
        if excerpts:
            matches.append({
                "filename": filename,
                "excerpts": excerpts[:5],  # cap at 5 lines per document
            })

    return {
        "status":      "success",
        "tag":         tag,
        "match_count": sum(len(m["excerpts"]) for m in matches),
        "matches":     matches,
    }


def extract_equipment_tags_from_text(text: str) -> List[str]:
    """
    Extracts unique equipment tags from arbitrary text using the standard regex pattern.
    Used by the planner to auto-detect tags after scan_ingest.
    """
    return sorted(set(_EQUIP_TAG_RE.findall(text)))


if __name__ == "__main__":
    result = lookup_equipment_history("CP-402-B")
    import json
    print(json.dumps(result, indent=2))
